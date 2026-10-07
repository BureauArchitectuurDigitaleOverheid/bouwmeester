"""A Word document is read whole, tables included.

A "Lijst van vragen" of the Tweede Kamer is one table, from the first
question to the last. Reading the loose paragraphs only gave the cover
page, and the model then judged a cover page.
"""

from pathlib import Path

from docx import Document

from bouwmeester.services.document_extract import extract_text

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _write(tmp_path: Path, build) -> Path:
    doc = Document()
    build(doc)
    path = tmp_path / "stuk.docx"
    doc.save(str(path))
    return path


def _lijst_van_vragen(doc) -> None:
    doc.add_paragraph("Verslag houdende een lijst van vragen")
    doc.add_paragraph("De commissie heeft de volgende vragen voorgelegd.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Nr"
    table.cell(0, 1).text = "Vraag"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "Welke middelen zijn er voor de voorbeelddienst?"
    table.cell(2, 0).text = "2"
    table.cell(2, 1).text = "Wanneer is de voorbeeldcloud beschikbaar?"
    doc.add_paragraph("De voorzitter van de commissie")


class TestTabellen:
    def test_the_questions_in_a_table_are_read(self, tmp_path):
        """The case from production: the terms stood in the table only."""
        tekst = extract_text(_write(tmp_path, _lijst_van_vragen), DOCX)

        assert "voorbeelddienst" in tekst
        assert "voorbeeldcloud" in tekst

    def test_in_the_order_of_the_document(self, tmp_path):
        tekst = extract_text(_write(tmp_path, _lijst_van_vragen), DOCX)

        assert tekst.splitlines() == [
            "Verslag houdende een lijst van vragen",
            "De commissie heeft de volgende vragen voorgelegd.",
            "Nr | Vraag",
            "1 | Welke middelen zijn er voor de voorbeelddienst?",
            "2 | Wanneer is de voorbeeldcloud beschikbaar?",
            "De voorzitter van de commissie",
        ]

    def test_a_document_without_a_table_reads_as_before(self, tmp_path):
        def build(doc):
            doc.add_paragraph("Eerste alinea.")
            doc.add_paragraph("   ")
            doc.add_paragraph("Tweede alinea.")

        assert extract_text(_write(tmp_path, build), DOCX) == (
            "Eerste alinea.\nTweede alinea."
        )

    def test_a_merged_cell_is_read_once(self, tmp_path):
        def build(doc):
            table = doc.add_table(rows=2, cols=3)
            table.cell(0, 0).merge(table.cell(0, 2)).text = "Kop over drie kolommen"
            table.cell(1, 0).text = "a"
            table.cell(1, 1).text = "b"
            table.cell(1, 2).text = "c"

        tekst = extract_text(_write(tmp_path, build), DOCX)

        assert tekst.splitlines() == ["Kop over drie kolommen", "a | b | c"]

    def test_empty_cells_and_rows_leave_nothing(self, tmp_path):
        def build(doc):
            table = doc.add_table(rows=3, cols=2)
            table.cell(0, 1).text = "alleen rechts"
            table.cell(2, 0).text = "alleen links"

        tekst = extract_text(_write(tmp_path, build), DOCX)

        assert tekst.splitlines() == ["alleen rechts", "alleen links"]

    def test_a_table_in_a_cell_is_read_too(self, tmp_path):
        def build(doc):
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "buiten"
            binnen = table.cell(0, 1).add_table(rows=1, cols=2)
            binnen.cell(0, 0).text = "diep"
            binnen.cell(0, 1).text = "dieper"

        tekst = extract_text(_write(tmp_path, build), DOCX)

        assert "buiten" in tekst
        assert "diep | dieper" in tekst

    def test_several_paragraphs_in_one_cell_stay_on_the_row(self, tmp_path):
        def build(doc):
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "7"
            cel = table.cell(0, 1)
            cel.text = "Eerste zin van de vraag."
            cel.add_paragraph("Tweede zin van de vraag.")

        tekst = extract_text(_write(tmp_path, build), DOCX)

        assert tekst == "7 | Eerste zin van de vraag. Tweede zin van de vraag."

    def test_the_limit_on_length_still_holds(self, tmp_path):
        def build(doc):
            table = doc.add_table(rows=50, cols=1)
            for n in range(50):
                table.cell(n, 0).text = f"Vraag {n}: " + "woord " * 40

        tekst = extract_text(_write(tmp_path, build), DOCX, max_chars=500)

        assert len(tekst) <= 600
        assert tekst.startswith("Vraag 0:")
